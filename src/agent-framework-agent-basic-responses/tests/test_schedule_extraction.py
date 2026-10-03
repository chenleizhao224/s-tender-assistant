import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from tender_assistant.documents.inventory import sha256_file
from tender_assistant.schedules.extraction import SHEET_NAME, extract_schedule


class ScheduleExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "schedule.xlsx"

    def make_workbook(self, rows, *, hidden_rows=()):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME

        for address, value in {
            "B6": "Code",
            "C6": "Description",
            "D6": "Quantity",
            "E6": "Unit",
        }.items():
            sheet[address] = value

        for row_number, values in enumerate(rows, start=7):
            for column, value in enumerate(values, start=2):
                sheet.cell(row_number, column, value)

        for row_number in hidden_rows:
            sheet.row_dimensions[row_number].hidden = True

        workbook.save(self.path)
        workbook.close()

    def test_numeric_item_preserves_source(self):
        self.make_workbook([
            ("5.1", "Cut to waste", 95, "m3"),
        ])

        result = extract_schedule(self.path)
        row = result["rows"][0]
        fingerprint = sha256_file(self.path)

        self.assertEqual(result["source_file"], "schedule.xlsx")
        self.assertEqual(result["source_sha256"], fingerprint)
        self.assertEqual(result["sheet"], SHEET_NAME)
        self.assertEqual(result["columns_read"], ["B", "C", "D", "E"])
        self.assertEqual(
            row["row_id"],
            f"{fingerprint}:{SHEET_NAME}:7",
        )
        self.assertEqual(row["excel_row"], 7)
        self.assertEqual(row["fields"]["quantity"]["value"], "95")
        self.assertEqual(row["fields"]["quantity"]["coordinate"], "D7")
        self.assertEqual(row["fields"]["quantity"]["data_type"], "n")
        self.assertEqual(row["issues"], [])

    def test_formula_is_preserved_not_evaluated(self):
        self.make_workbook([
            ("5.18", "Trench backfill", "=25*0.6*0.55", "m3"),
        ])

        row = extract_schedule(self.path)["rows"][0]

        self.assertEqual(
            row["fields"]["quantity"]["value"],
            "=25*0.6*0.55",
        )
        self.assertEqual(row["fields"]["quantity"]["data_type"], "f")
        self.assertIn(
            "Quantity is a formula; not evaluated.",
            row["issues"],
        )

    def test_special_text_is_preserved(self):
        self.make_workbook([
            (None, "Scope explanation", "Note", None),
            ("8.1", "Contractor margins", "Incl.", "%"),
            ("9.17", "Other plant", "Rate Only", "hr"),
        ])

        rows = extract_schedule(self.path)["rows"]

        self.assertEqual(
            [row["fields"]["quantity"]["value"] for row in rows],
            ["Note", "Incl.", "Rate Only"],
        )
        for row in rows:
            self.assertIn(
                "Quantity is not a numeric Excel cell.",
                row["issues"],
            )

    def test_missing_quantity_stays_none(self):
        self.make_workbook([
            ("1.12", "On-site overheads", None, "%"),
        ])

        row = extract_schedule(self.path)["rows"][0]

        self.assertIsNone(row["fields"]["quantity"]["value"])
        self.assertIn(
            "Unit supplied without a quantity.",
            row["issues"],
        )

    def test_zero_quantity_is_not_missing(self):
        self.make_workbook([
            ("5.1", "Cut to waste", 0, "m3"),
        ])

        row = extract_schedule(self.path)["rows"][0]

        self.assertEqual(row["fields"]["quantity"]["value"], "0")
        self.assertEqual(row["issues"], [])

    def test_blank_rows_skipped_but_headings_retained(self):
        self.make_workbook([
            (None, "EARTHWORKS", None, None),
            (None, None, None, None),
            ("5.1", "Cut to waste", 95, "m3"),
        ])

        rows = extract_schedule(self.path)["rows"]

        self.assertEqual([row["excel_row"] for row in rows], [7, 9])
        self.assertEqual(
            rows[0]["fields"]["description"]["value"],
            "EARTHWORKS",
        )

    def test_hidden_rows_are_retained_and_flagged(self):
        self.make_workbook(
            [("5.1", "Cut to waste", 95, "m3")],
            hidden_rows=(7,),
        )

        rows = extract_schedule(self.path)["rows"]

        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["hidden"])

    def test_missing_unit_and_description_are_flagged(self):
        self.make_workbook([
            ("5.1", None, 95, None),
        ])

        row = extract_schedule(self.path)["rows"][0]

        self.assertIn(
            "Quantity supplied without a unit.",
            row["issues"],
        )
        self.assertIn(
            "Possible item has no description.",
            row["issues"],
        )

    def test_missing_worksheet_is_rejected(self):
        workbook = Workbook()
        workbook.save(self.path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "Missing worksheet"):
            extract_schedule(self.path)

    def test_unexpected_header_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet["B6"] = "Wrong header"
        workbook.save(self.path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "Unexpected header"):
            extract_schedule(self.path)

    def test_repeated_extraction_is_identical(self):
        self.make_workbook([
            ("5.1", "Cut to waste", 95, "m3"),
        ])

        self.assertEqual(
            extract_schedule(self.path),
            extract_schedule(self.path),
        )


if __name__ == "__main__":
    unittest.main()
