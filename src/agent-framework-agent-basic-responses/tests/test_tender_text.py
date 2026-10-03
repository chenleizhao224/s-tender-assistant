import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfWriter

from tender_assistant.text.extraction import extract_document, find_evidence


class TenderTextTests(unittest.TestCase):
    def test_content_controls_and_table_cells_are_not_lost(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "controlled.docx"
            doc = Document()
            paragraph = doc.add_paragraph("Imported fill requires approval.")
            control = OxmlElement("w:sdt")
            content = OxmlElement("w:sdtContent")
            control.append(content)
            paragraph._p.addprevious(control)
            content.append(paragraph._p)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Compaction requirement"
            doc.save(path)
            extracted = extract_document(path)
            texts = [segment["text"] for segment in extracted["segments"]]
            self.assertEqual(texts, ["Imported fill requires approval.", "Compaction requirement"])
            self.assertTrue(all("xml_path" in segment for segment in extracted["segments"]))
            self.assertEqual(extracted, extract_document(path))

    def test_blank_pdf_page_requires_review(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            writer.write(path)
            result = extract_document(path)
            self.assertEqual(result["status"], "partial_text")
            self.assertEqual(result["segments"][0]["locator"], "pdf-page-1")
            self.assertEqual(result["segments"][0]["status"], "sparse_text_review_required")

    def test_keyword_context_is_ordered_and_deduplicated(self):
        doc = {"source_file": "test", "source_sha256": "abc", "segments": [
            {"text": text, "locator": str(index)}
            for index, text in enumerate(["before", "fill", "fill compaction", "after", "unrelated"])
        ]}
        result = find_evidence([doc], ["fill"])
        self.assertEqual([segment["locator"] for segment in result], ["0", "1", "2", "3"])
        self.assertEqual(find_evidence([doc], ["not present"]), [])


if __name__ == "__main__":
    unittest.main()
