from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from document_inventory import sha256_file


SHEET_NAME = "Civil  SoP - Breakdown"


def cell_value(cell):
    """Keep values serializable, preserving formulas and cell locations."""
    value = cell.value

    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif value is not None:
        value = str(value)

    return {
        "value": value,
        "coordinate": cell.coordinate,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
    }


def extract_schedule(path: str | Path) -> dict:
    path = Path(path)
    fingerprint = sha256_file(path)

    workbook = load_workbook(
        path,
        read_only=False,
        data_only=False,
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet: {SHEET_NAME}")

        sheet = workbook[SHEET_NAME]

        expected_headers = {
            "B6": "Code",
            "C6": "Description",
            "D6": "Quantity",
            "E6": "Unit",
        }

        for address, expected in expected_headers.items():
            actual = str(sheet[address].value or "").strip()
            if actual.casefold() != expected.casefold():
                raise ValueError(
                    f"Unexpected header at {address}: {actual!r}"
                )

        rows = []

        for row_number in range(7, sheet.max_row + 1):
            cells = [
                sheet.cell(row=row_number, column=column)
                for column in range(2, 6)
            ]

            if all(
                cell.value is None
                or str(cell.value).strip() == ""
                for cell in cells
            ):
                continue

            fields = dict(zip(
                ("code", "description", "quantity", "unit"),
                map(cell_value, cells),
            ))

            issues = []
            quantity = fields["quantity"]
            unit = fields["unit"]["value"]
            description = fields["description"]["value"]

            has_quantity = bool(
                quantity["value"] and quantity["value"].strip()
            )
            has_unit = bool(unit and unit.strip())

            if quantity["data_type"] == "f":
                issues.append("Quantity is a formula; not evaluated.")
            elif has_quantity and quantity["data_type"] != "n":
                issues.append("Quantity is not a numeric Excel cell.")

            if has_quantity and not has_unit:
                issues.append("Quantity supplied without a unit.")

            if has_unit and not has_quantity:
                issues.append("Unit supplied without a quantity.")

            if (has_quantity or has_unit) and not (
                description and description.strip()
            ):
                issues.append("Possible item has no description.")

            rows.append({
                "row_id": (
                    f"{fingerprint}:{SHEET_NAME}:{row_number}"
                ),
                "excel_row": row_number,
                "hidden": bool(sheet.row_dimensions[row_number].hidden),
                "fields": fields,
                "issues": issues,
            })

        return {
            "source_file": path.name,
            "source_sha256": fingerprint,
            "sheet": SHEET_NAME,
            "columns_read": ["B", "C", "D", "E"],
            "rows": rows,
        }
    finally:
        workbook.close()