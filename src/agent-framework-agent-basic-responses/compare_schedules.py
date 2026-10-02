from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/chenlei/Downloads/01 Tender Documents")
SHEET = "Civil  SoP - Breakdown"

OLD = ROOT / "NTT 01" / (
    "20260701 - Hopuhopu Development - Schedule of Prices [0].xlsx"
)
NEW = ROOT / "NTT 02" / (
    "20260707 - Hopuhopu Development - Schedule of Prices [0A].xlsx"
)


def read_rows(path):
    workbook = load_workbook(
        path, read_only=True, data_only=False
    )

    try:
        sheet = workbook[SHEET]
        rows = []

        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=7,
                min_col=2,
                max_col=5,
                values_only=True,
            ),
            start=7,
        ):
            # Keep descriptions and headings, even without quantities.
            if all(
                value is None or str(value).strip() == ""
                for value in values
            ):
                continue

            # Preserve text and distinguish text from numeric cells.
            signature = tuple(
                (type(value).__name__, str(value))
                for value in values
            )
            rows.append((row_number, values, signature))

        return rows
    finally:
        workbook.close()


def show_rows(label, rows):
    for row_number, values, _ in rows:
        code, description, quantity, unit = values
        print(
            f"{label} row {row_number}: "
            f"code={code!r} | description={description!r} | "
            f"quantity={quantity!r} | unit={unit!r}"
        )


if __name__ == "__main__":
    old_rows = read_rows(OLD)
    new_rows = read_rows(NEW)

    matcher = SequenceMatcher(
        a=[row[2] for row in old_rows],
        b=[row[2] for row in new_rows],
        autojunk=False,
    )

    changes = 0

    for action, old_start, old_end, new_start, new_end in (
        matcher.get_opcodes()
    ):
        if action == "equal":
            continue

        changes += 1
        print(f"\nCHANGE {changes}: {action}")
        show_rows("OLD", old_rows[old_start:old_end])
        show_rows("NEW", new_rows[new_start:new_end])

    print(f"\nChanged blocks: {changes}")
    print("Compared Breakdown columns B–E only.")