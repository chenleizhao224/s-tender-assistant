from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/chenlei/Downloads/01 Tender Documents")


def inspect_workbook(path: Path) -> None:
    print(f"\nFILE: {path.relative_to(ROOT)}")

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
    )

    try:
        for sheet in workbook.worksheets:
            print(
                f"\nSHEET: {sheet.title!r}"
                f" | rows={sheet.max_row}"
                f" | columns={sheet.max_column}"
                f" | state={sheet.sheet_state}"
            )

            shown = 0

            for row in sheet.iter_rows():
                populated = [
                    cell
                    for cell in row
                    if cell.value is not None
                ]
                if not populated:
                    continue

                # Preview the first 8 populated rows of each sheet.
                for cell in populated[:12]:
                    text = str(cell.value).replace("\n", " ")
                    print(f"  {cell.coordinate}: {text[:180]}")

                shown += 1
                if shown >= 8:
                    break
    finally:
        workbook.close()


if __name__ == "__main__":
    paths = sorted(ROOT.rglob("*.xlsx"))

    if not paths:
        raise SystemExit("No Excel files found.")

    for path in paths:
        inspect_workbook(path)